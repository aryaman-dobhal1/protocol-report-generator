"""Smoke test for the three mandatory demo cases. Run from backend/:  python tests/test_cases.py
Expects the test-inputs folder next to backend/ (project root)."""
import io, os, sys, zipfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from app.main import app

T = os.environ.get("TEST_INPUTS", os.path.join(os.path.dirname(__file__), "..", "..", "test-inputs"))
client = TestClient(app)


def files(case, with_pdf=True):
    sh = os.path.join(T, "shared")
    f = [("protocol", ("protocol.json", open(f"{sh}/protocol.json", "rb"))),
         ("results", ("results.csv", open(f"{T}/{case}/results.csv", "rb"))),
         ("template", ("report-template.docx", open(f"{sh}/report-template.docx", "rb")))]
    f += [("references", (n, open(f"{sh}/reference/{n}", "rb"))) for n in sorted(os.listdir(f"{sh}/reference"))]
    if with_pdf:
        f.append(("calibration", ("calibration-record.pdf", open(f"{sh}/calibration-record.pdf", "rb"))))
    return f


def run(case, with_pdf=True):
    j = client.post("/api/analyze", files=files(case, with_pdf)).json()
    z = client.post("/api/export", files=files(case, with_pdf), data={"format": "zip", "summary": j["summary"]})
    assert z.status_code == 200, z.text
    zf = zipfile.ZipFile(io.BytesIO(z.content))
    docx = Document(io.BytesIO(zf.read("Report.docx")))
    text = "\n".join(p.text for p in docx.paragraphs) + "\n".join(c.text for t in docx.tables for r in t.rows for c in r.cells)
    ws = load_workbook(io.BytesIO(zf.read("Attachment B - Results.xlsx")))["Results"]
    xl = {r[0].value: r[6].value for r in ws.iter_rows(min_row=2, max_row=1 + j["total"])}
    pdf = PdfReader(io.BytesIO(zf.read("Report_combined.pdf")))
    return j, zf, text, xl, pdf


j, zf, text, xl, pdf = run("case-1-complete")
assert (j["total"], j["passed"], j["failed"]) == (5, 5, 0) and all(c["status"] == "ok" for c in j["checks"])
assert "{{" not in text and set(xl.values()) == {"Pass"} and len(zf.namelist()) == 4
print("case 1 ok:", sorted(zf.namelist()), "pdf pages", len(pdf.pages))

j, zf, text, xl, pdf = run("case-2-failing")
assert j["failed"] == 1 and xl["S-003"] == "Fail" and "2.71" in text and "S-003" in j["summary"] and "Fail" in text
print("case 2 ok: S-003 Fail in docx, xlsx, summary;", [c["detail"] for c in j["checks"] if c["status"] == "bad"])

j, zf, text, xl, pdf = run("case-3-incomplete", with_pdf=False)
assert j["missing"] == ["S-004"] and "{{calibration_ref}}" in j["unfilled"] and "{{calibration_ref}}" in text
assert len(zf.namelist()) == 3 and "not supplied" in "".join(p.extract_text() for p in pdf.pages)
assert "S-004" in j["summary"] and "calibration" in j["summary"].lower()
print("case 3 ok:", [c["title"] for c in j["checks"] if c["status"] != "ok"])
print("all cases passed")